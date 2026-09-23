"""Translator-afterword generation, reuse and export tests."""

from __future__ import annotations

import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as open_docx
from wenyi_core.afterword import (
    AFTERWORD_ARTIFACT,
    AFTERWORD_DRAFT_ARTIFACT,
    afterword_content_digest,
)
from wenyi_core.agents.afterword import AfterwordWriter
from wenyi_core.assemble.writer import assemble
from wenyi_core.config import Config
from wenyi_core.ingest.models import Chapter, Document, Segment
from wenyi_core.llm.providers.fake import FakeClient
from wenyi_core.pipeline.afterword import AfterwordService
from wenyi_core.pipeline.runstore import source_sha256
from wenyi_core.pipeline.runtime import PipelineRuntime
from wenyi_core.storage.file import FileStorage
from wenyi_core.storage.protocol import STATUS_DONE


def _afterword_response(label: str) -> str:
    return json.dumps(
        {
            "title": "译者后记",
            "paragraphs": [
                f"{label}：作者以克制的笔法处理这一题材。",
                "作品的写作背景使其关注个人选择与制度压力。",
                "它的长处在于结构逐步收紧，人物行动能够推动主题。",
                "局限则是部分次要人物较为功能化，中段节奏也略显重复。",
                "翻译保留了关键称谓的差异，以呈现人物关系的变化。",
                "这些得失共同构成了作品值得重新阅读的复杂面貌。",
            ],
        },
        ensure_ascii=False,
    )


def _handler(messages, _tier, _json_mode):
    if "final critical editor" in messages[0]["content"]:
        return _afterword_response("修订稿")
    if "translator's afterword" in messages[0]["content"]:
        return _afterword_response("草稿")
    return "{}"


def _completed_store(directory: str) -> tuple[FileStorage, str]:
    source = os.path.join(directory, "book.txt")
    Path(source).write_text("source", encoding="utf-8")
    chapter = Chapter(
        index=0,
        title="第一章",
        segments=[Segment(index=0, source="原文", target="译文")],
        meta={"source_digest": "主人公在限制中作出选择。"},
    )
    document = Document(
        title="测试作品",
        source_lang="ja",
        target_lang="zh",
        fmt="text",
        source_path=source,
        chapters=[chapter],
        meta={"authors": ["测试作者"]},
    )
    store = FileStorage(os.path.join(directory, "state"))
    store.init_from_document(document)
    store.set_chapter_status(0, STATUS_DONE)
    store.save_analysis(
        {
            "tone": "克制",
            "style_guide": "避免夸饰",
            "book_synopsis": "主人公在制度约束下重新理解责任。",
        }
    )
    return store, source


def _saved_afterword(store: FileStorage, label: str) -> dict:
    return {
        "version": 1,
        "fingerprint": {"content_digest": afterword_content_digest([store.load_chapter(0)])},
        "content": json.loads(_afterword_response(label)),
    }


class TestAfterwordAgent(unittest.TestCase):
    def test_draft_is_critically_revised_with_grounded_prompt(self):
        config = Config.from_dict(
            {
                "language": {"source": "ja", "target": "zh"},
                "llm": {"preset": "fake"},
            }
        )
        client = FakeClient(handler=_handler)
        writer = AfterwordWriter(client, config)

        result = writer.generate(
            book_title="测试作品",
            authors=["测试作者"],
            verified_context="出版背景材料",
            book_synopsis="全书梗概",
            chapter_digests="章节摘要",
            representative_passages="原文和译文样段",
            style_brief="克制",
            glossary_terms=[],
        )

        self.assertEqual(result.title, "译者后记")
        self.assertTrue(result.paragraphs[0].startswith("修订稿"))
        self.assertEqual(
            [call["operation"] for call in client.calls],
            ["afterword.draft", "afterword.revise"],
        )
        self.assertIn("This is not promotional copy", client.calls[0]["messages"][0]["content"])
        self.assertIn("原文和译文样段", client.calls[0]["messages"][1]["content"])
        self.assertIn(
            "at least one evidence-based strength", client.calls[1]["messages"][0]["content"]
        )


class TestAfterwordPipeline(unittest.TestCase):
    def test_export_snapshot_keeps_the_saved_afterword(self):
        with tempfile.TemporaryDirectory() as directory:
            store, source = _completed_store(directory)
            store.write_artifact(AFTERWORD_ARTIFACT, _saved_afterword(store, "原稿"))
            snapshot = store.create_export_snapshot(actual_sha256=source_sha256(source))
            store.write_artifact(AFTERWORD_ARTIFACT, _saved_afterword(store, "新稿"))

            path = os.path.join(directory, "snapshot.txt")
            assemble(snapshot, source, out_path=path, out_format="txt")
            content = Path(path).read_text(encoding="utf-8")
            self.assertIn("原稿：作者", content)
            self.assertNotIn("新稿：作者", content)

    def test_completed_fingerprint_is_reused_and_translation_change_invalidates_it(self):
        with tempfile.TemporaryDirectory() as directory:
            store, _source = _completed_store(directory)
            config = Config.from_dict(
                {
                    "language": {"source": "ja", "target": "zh"},
                    "llm": {"preset": "fake"},
                    "pipeline": {
                        "translator_afterword": True,
                        "translator_afterword_context": "可靠背景",
                    },
                }
            )
            client = FakeClient(handler=_handler)
            runtime = PipelineRuntime(config, client=client)
            runtime.apply_language("ja")
            service = AfterwordService(runtime)

            first = service.generate(store)
            second = service.generate(store)

            self.assertEqual(first, second)
            self.assertEqual(len(client.calls), 2)
            self.assertIn("Source: 原文", client.calls[0]["messages"][1]["content"])
            self.assertIn("Translation: 译文", client.calls[0]["messages"][1]["content"])
            self.assertIsNotNone(store.read_artifact(AFTERWORD_ARTIFACT))

            chapter = store.load_chapter(0)
            chapter.segments[0].target = "修改后的译文"
            store.save_chapter(chapter)
            service.generate(store)
            self.assertEqual(len(client.calls), 4)

            chapter = store.load_chapter(0)
            chapter.meta["source_digest"] = "补充后的章节摘要"
            store.save_chapter(chapter)
            service.generate(store)
            self.assertEqual(len(client.calls), 6)

            runtime.afterword_writer.tgt = "en"
            service.generate(store)
            self.assertEqual(len(client.calls), 8)

    def test_revision_failure_reuses_saved_draft(self):
        with tempfile.TemporaryDirectory() as directory:
            store, _source = _completed_store(directory)
            config = Config.from_dict(
                {"language": {"source": "ja", "target": "zh"}, "llm": {"preset": "fake"}}
            )
            attempts = 0

            def fail_once(messages, tier, json_mode):
                nonlocal attempts
                if "final critical editor" in messages[0]["content"]:
                    attempts += 1
                    if attempts == 1:
                        raise RuntimeError("revision unavailable")
                return _handler(messages, tier, json_mode)

            client = FakeClient(handler=fail_once)
            runtime = PipelineRuntime(config, client=client)
            runtime.apply_language("ja")
            service = AfterwordService(runtime)
            with self.assertRaisesRegex(RuntimeError, "revision unavailable"):
                service.generate(store)
            self.assertIsNotNone(store.read_artifact(AFTERWORD_DRAFT_ARTIFACT))
            self.assertIsNone(store.read_artifact(AFTERWORD_ARTIFACT))

            service.generate(store)
            self.assertEqual(
                [call["operation"] for call in client.calls],
                ["afterword.draft", "afterword.revise", "afterword.revise"],
            )
            self.assertIsNotNone(store.read_artifact(AFTERWORD_ARTIFACT))

    def test_stale_afterword_requires_regeneration_or_explicit_omission(self):
        with tempfile.TemporaryDirectory() as directory:
            store, source = _completed_store(directory)
            store.write_artifact(AFTERWORD_ARTIFACT, _saved_afterword(store, "旧稿"))
            chapter = store.load_chapter(0)
            chapter.segments[0].target = "改过的译文"
            store.save_chapter(chapter)

            with self.assertRaisesRegex(ValueError, "out of date"):
                assemble(store, source, out_format="txt")
            path = assemble(store, source, out_format="txt", include_translator_afterword=False)
            self.assertNotIn("旧稿", Path(path).read_text(encoding="utf-8"))

            chapter.segments[0].target = "译文"
            chapter.title = "改过的章名"
            store.save_chapter(chapter)
            with self.assertRaisesRegex(ValueError, "out of date"):
                assemble(store, source, out_format="txt")

    def test_afterword_is_exported_once_and_precedes_about_page(self):
        with tempfile.TemporaryDirectory() as directory:
            store, source = _completed_store(directory)
            store.write_artifact(AFTERWORD_ARTIFACT, _saved_afterword(store, "定稿"))

            text_path = assemble(store, source, out_format="txt", bilingual=True)
            text = Path(text_path).read_text(encoding="utf-8")
            self.assertEqual(text.count("译者后记"), 1)
            self.assertIn("部分次要人物较为功能化", text)

            html_path = assemble(store, source, out_format="html")
            html = BeautifulSoup(Path(html_path).read_text(encoding="utf-8"), "html.parser")
            section = html.select_one("section.wenyi-translator-afterword")
            self.assertIsNotNone(section)
            self.assertEqual(section.h1.get_text(), "译者后记")

            docx_path = assemble(store, source, out_format="docx")
            docx = open_docx(docx_path)
            self.assertEqual(sum(p.text == "译者后记" for p in docx.paragraphs), 1)

            epub_path = assemble(store, source, out_format="epub")
            with zipfile.ZipFile(epub_path) as archive:
                afterword_name = next(
                    name
                    for name in archive.namelist()
                    if name.endswith("wenyi-translator-afterword.xhtml")
                )
                self.assertIn("部分次要人物较为功能化", archive.read(afterword_name).decode())
                opf_name = next(name for name in archive.namelist() if name.endswith("content.opf"))
                package = BeautifulSoup(archive.read(opf_name), "xml")
                spine_ids = [item.get("idref") for item in package.find_all("itemref")]
                self.assertEqual(
                    spine_ids[-2:], ["wenyi-translator-afterword", "trans-novel-about"]
                )
                nav_name = next(name for name in archive.namelist() if name.endswith("nav.xhtml"))
                self.assertIn("译者后记", archive.read(nav_name).decode("utf-8"))

            without_path = os.path.join(directory, "without.txt")
            assemble(
                store,
                source,
                out_path=without_path,
                out_format="txt",
                include_translator_afterword=False,
            )
            self.assertNotIn("译者后记", Path(without_path).read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
