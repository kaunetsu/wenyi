"""Create the DOCX document, coordinate chapter emission and save the export."""

from __future__ import annotations

from docx import Document as open_docx
from docx.oxml.ns import qn

from wenyi_core.afterword import TranslatorAfterword
from wenyi_core.assemble.docx_blocks import _add_heading, _add_normal, _emit_chapter_blocks
from wenyi_core.assemble.docx_styles import _target_output_font
from wenyi_core.assemble.writer_common import (
    _ch_title,
    _manifest_target_lang,
)
from wenyi_core.ingest.models import KIND_HEADING

from .export_view import AssembleStore


def _assemble_docx(
    store: AssembleStore,
    out_path: str,
    *,
    bilingual: bool = False,
    order: str = "target_first",
    translator_afterword: TranslatorAfterword | None = None,
) -> str:
    """Rebuild a DOCX by chapter, restoring heading outlines, styles and tables from metadata."""
    manifest = store.load_manifest()
    output_font = _target_output_font(_manifest_target_lang(manifest))
    doc = open_docx()
    if doc.paragraphs:
        p0 = doc.paragraphs[0]
        if not p0.text.strip():
            p0.clear()

    first_block = True
    for c in manifest["chapters"]:
        chapter = store.load_chapter(c["index"])
        has_h1 = any(
            s.kind == KIND_HEADING
            and isinstance(s.meta, dict)
            and int(s.meta.get("heading_level") or 1) == 1
            for s in chapter.segments
        )
        title = _ch_title(c)
        if title and not has_h1 and chapter.meta.get("explicit_title"):
            if first_block and doc.paragraphs and not doc.paragraphs[0].text:
                pass
            _add_heading(doc, title, 1, output_font=output_font)
        _emit_chapter_blocks(
            doc,
            chapter,
            bilingual=bilingual,
            order=order,
            output_font=output_font,
        )
        first_block = False

    if translator_afterword is not None:
        doc.add_page_break()
        _add_heading(doc, translator_afterword.title, 1, output_font=output_font)
        for paragraph in translator_afterword.paragraphs:
            _add_normal(doc, paragraph, output_font=output_font)

    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p") and not (child.text or "").strip():
            texts = [node.text for node in child.iter(qn("w:t")) if node.text]
            if not any(texts) and body.index(child) == 0:
                body.remove(child)
            break

    doc.save(out_path)
    return out_path
